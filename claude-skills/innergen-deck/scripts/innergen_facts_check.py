#!/usr/bin/env python3
"""
INNerGEN company-facts gate - compares a deck with the company facts ledger
(references/company-facts-v28.json) so recurring numbers, contacts and brand
spellings stay current across every deck.

Codes:
    F01  number or contact differs from the ledger (stale value)        WARN
    F02  brand spelling mixed within the deck, or not the ledger form   WARN
    F03  a value the ledger marks needs_confirm is used                 INFO
    F04  a ledger entry is past its review_by date                      INFO

Usage:
    python3 scripts/innergen_facts_check.py --pptx deck.pptx       # existing PowerPoint (needs python-pptx)
    python3 scripts/innergen_facts_check.py --svg /tmp/<deck>      # svg_output/*.svg of a skill project
    python3 scripts/innergen_facts_check.py --outline outline.md   # page plan
    python3 scripts/innergen_facts_check.py --print-md -o references/company-facts-v28.md   # ledger as a fact sheet

innergen_qaqc.py calls check() on every finished deck. To change a fact, edit
company-facts-v28.json (value, as_of, source, and the matching `checks` entry).

Dependencies:
    Standard library; python-pptx only for --pptx
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

_SCRIPTS_DIR = Path(__file__).resolve().parent
REPO = _SCRIPTS_DIR.parents[min(3, len(_SCRIPTS_DIR.parents) - 1)]
FACTS_PATH = REPO / "docs/styles/innergen/company-facts-v28.json"
# Standalone skill bundles ship the ledger beside the skill.
if not FACTS_PATH.is_file():
    FACTS_PATH = _SCRIPTS_DIR.parent / "references/company-facts-v28.json"


def load_facts(path: Path | None = None) -> dict | None:
    p = path or FACTS_PATH
    if not p.is_file():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


def _digits(s: str) -> str:
    return re.sub(r"\D", "", s)


def _walk(node, path=""):
    """Yield (dotted path, dict) for every ledger entry that has a status."""
    if isinstance(node, dict):
        if "status" in node and ("value" in node or "claim" in node):
            yield path, node
        for k, v in node.items():
            if isinstance(v, (dict, list)):
                yield from _walk(v, f"{path}.{k}" if path else k)
    elif isinstance(node, list):
        for i, v in enumerate(node):
            key = v.get("id", str(i)) if isinstance(v, dict) else str(i)
            yield from _walk(v, f"{path}.{key}")


def check(pages: list[tuple[int, list[str]]], facts: dict, add, today: _dt.date | None = None) -> dict:
    """pages: [(page_no, [text, ...])]; add(level, code, page, message)."""
    today = today or _dt.date.today()
    found = 0

    for c in facts.get("checks", []):
        rx = re.compile(c["rx"])
        page_rx = re.compile(c["page_rx"]) if c.get("page_rx") else None
        norm = _digits if c.get("normalize") == "digits" else (lambda s: s.strip())
        expect = {norm(e) for e in c["expect"]}
        for no, texts in pages:
            blob = "\n".join(texts)
            if page_rx and not page_rx.search(blob):
                continue
            for m in rx.finditer(blob):
                got = next((g for g in m.groups() if g), m.group(0))
                if norm(got) not in expect:
                    found += 1
                    add("WARN", "F01", no, f"{c['label']}: 장표 '{got}' ≠ 원장 '{c['expect'][0]}' "
                        f"(company-facts {c.get('fact', c['id'])})")

    brand = facts.get("brand", {})
    for key, pattern in facts.get("brand_rx", {}).items():
        entry = brand.get(f"{key}_spelling", {})
        canon = entry.get("value")
        forms: dict[str, int] = {}
        for no, texts in pages:
            for m in re.finditer(pattern, "\n".join(texts), re.IGNORECASE):
                form = re.sub(r"\s", "", m.group(0))
                if form.islower():  # e-mail and web addresses
                    continue
                forms.setdefault(form, no)
        if len(forms) > 1:
            add("WARN", "F02", 0, f"브랜드 표기가 섞여 있습니다: {', '.join(forms)} — 한 가지로 통일 "
                f"(원장 {canon}, 상태 {entry.get('status', '-')})")
        elif forms and canon and entry.get("status") == "confirmed" and canon not in forms:
            form, no = next(iter(forms.items()))
            add("WARN", "F02", no, f"브랜드 표기 '{form}' ≠ 원장 '{canon}'")

    blob_all = "\n".join(t for _, ts in pages for t in ts)
    for path, entry in _walk(facts):
        val = entry.get("value")
        if entry.get("status") == "needs_confirm" and isinstance(val, str) and len(val) <= 40 \
                and re.sub(r"\s", "", val) in re.sub(r"\s", "", blob_all):
            add("INFO", "F03", 0, f"확인이 필요한 원장 값을 썼습니다: {path} '{val}' — {entry.get('conflict') or entry.get('note') or entry.get('source', '')}")
        rb = entry.get("review_by")
        if rb:
            try:
                due = _dt.date.fromisoformat(rb)
            except ValueError:
                continue
            if due < today:
                add("INFO", "F04", 0, f"원장 항목 검토 기한 경과: {path} (review_by {rb}) — 값과 출처를 다시 확인")
    return {"stale": found}


# ---------------------------------------------------------------- readers
def pages_from_pptx(path: Path) -> list[tuple[int, list[str]]]:
    from pptx import Presentation  # optional dependency

    def walk(shapes):
        for s in shapes:
            if s.shape_type == 6:
                yield from walk(s.shapes)
            else:
                yield s

    out = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        texts = []
        for s in walk(slide.shapes):
            if s.has_text_frame and s.text_frame.text.strip():
                texts.append(s.text_frame.text)
            if getattr(s, "has_table", False) and s.has_table:
                texts += [c.text for r in s.table.rows for c in r.cells if c.text.strip()]
        out.append((i, texts))
    return out


def pages_from_svg(project: Path) -> list[tuple[int, list[str]]]:
    folder = project / "svg_output" if (project / "svg_output").is_dir() else project
    out = []
    for i, f in enumerate(sorted(folder.glob("*.svg")), 1):
        root = ET.parse(f).getroot()
        texts = ["".join(t.itertext()) for t in root.iter() if t.tag.endswith("text")]
        out.append((i, [t for t in texts if t.strip()]))
    return out


def pages_from_outline(path: Path) -> list[tuple[int, list[str]]]:
    out = []
    for line in path.read_text(encoding="utf-8").splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if line.strip().startswith("|") and cells and re.fullmatch(r"P?\d{1,3}", cells[0]):
            out.append((int(cells[0].lstrip("P")), cells[1:]))
    return out


def pages_from_qaqc(pages) -> list[tuple[int, list[str]]]:
    """Adapt innergen_qaqc.Page objects."""
    return [(p.index, [t["text"] for t in p.texts if t["text"]]) for p in pages]


# ---------------------------------------------------------------- fact sheet
def to_markdown(facts: dict) -> str:
    mark = {"confirmed": "확정", "needs_confirm": "**확인 필요**", "time_sensitive": "**시의성·출처 확인**"}
    co, msg = facts["company"], facts["messages"]
    md = [f"# 이너젠 회사 기본 정보 원장 ({facts['version']}, {facts['updated']})", "",
          "> 이 문서는 `company-facts-v28.json`에서 만든다(`innergen_facts_check.py --print-md -o <이 파일>`). 값을 고칠 때는 JSON을 고치고 다시 만든다.",
          "> 장표에 넣는 회사 숫자·연락처·문구는 여기서만 가져온다. **확인 필요** 값은 쓰기 전에 사용자에게 묻는다.", "",
          "출처: " + " · ".join(f"`{k}` {v}" for k, v in facts["sources"].items()), "",
          "## 1. 회사 개요", "", "| 항목 | 값 | 상태 | 출처 |", "|---|---|---|---|"]
    labels = {"name_ko": "회사명", "name_en": "영문명", "legal_name": "법인명", "ceo": "대표자", "headcount": "인원",
              "founded": "설립일", "address": "주소", "phone": "전화", "email": "이메일", "website": "웹사이트",
              "main_clients": "주요 고객"}
    for k, lab in labels.items():
        e = co[k]
        v = e["value"] + (f" ({e['postcode']})" if e.get("postcode") else "")
        md.append(f"| {lab} | {v} | {mark[e['status']]} | {e['source']} |")
    md += ["", "부문: " + " / ".join(f"{d['name']}({d['since']}~) {d['desc']}" for d in co["divisions"]["value"]), "",
           "## 2. 고정 문구", "", "| 쓰는 곳 | 문구 | 출처 |", "|---|---|---|"]
    for k, e in msg.items():
        md.append(f"| {k} | {e['value']} | {e['source']} |")
    md += ["", "## 3. 사업 분야(5)", "", "| 분야 | 키워드 | content_check 사업 성격 |", "|---|---|---|"]
    for a in facts["business_areas"]["value"]:
        md.append(f"| {a['name']} | {a['tags']} | `{a['content_business_type']}` |")
    md += ["", "## 4. 핵심 숫자", "", "| 항목 | 값 | 세부 | 상태 | 출처·메모 |", "|---|---|---|---|---|"]
    for m in facts["metrics"]:
        md.append(f"| {m['label']} | {m['value']} | {m.get('detail', '')} | {mark[m['status']]} | {m['source']}{' — ' + m['conflict'] if m.get('conflict') else ''}{' — ' + m['note'] if m.get('note') else ''} |")
    md += ["", "## 5. 대외 자격·협력", ""] + [f"- {c}" for c in facts["credentials"]["value"]]
    md += ["", f"- {facts['credentials']['partner_logos_note']}", "",
           "## 6. CDP·환경공시 이력 (**확인 필요**)", "", "| 연도 | 고객 수 | 사건 |", "|---|---|---|"]
    for r in facts["cdp_track_record"]["value"]:
        md.append(f"| {r['year']} | {r.get('clients', '')} | {r.get('event', '')} |")
    md += ["", facts["cdp_track_record"]["note"], "", "## 7. 감축사업·배출권", ""]
    rx = facts["reduction_experience"]
    md += [f"- {t}" for t in rx["themes"]] + [f"- 시장: {rx['markets']}", f"- {rx['issuance']}", f"- {rx['methodology_count']}",
                                              f"- 주의: {rx['client_names_note']}", "", "## 8. 제품 LCA 자동화 Tool", "",
                                              f"- {facts['lca_tool']['value']}", f"- 구축: {facts['lca_tool']['references']}", "",
                                              "## 9. 이너캐리", ""]
    ic = facts["innercarry"]
    md += [f"- 고민→해법: {' / '.join(ic['pains'])}",
           "- 도입 방식: " + " / ".join(f"{d['name']}({d['role']}) {d['desc']}" for d in ic["deployment"]),
           f"- 5단계: {' → '.join(ic['steps'])}", f"- 보안 4축: {' · '.join(ic['security'])}",
           f"- ICFR형 통제 5가지: {' · '.join(ic['controls'])}", f"- 도입 효과 5가지: {' · '.join(ic['benefits'])}",
           f"- 약속 5가지: {' · '.join(ic['promises'])}", f"- 표준: {ic['standards']}", f"- 범위: {ic['coverage']}",
           f"- 출시: {ic['launch']['value']} ({mark[ic['launch']['status']]}) — {ic['launch']['note']}", "",
           "## 10. 시의성 있는 외부 사실", ""]
    for t in facts["time_sensitive"]:
        md += [f"- {t['claim']} ({mark[t['status']]}, {t['source']}, 검토 기한 {t['review_by']})", f"  - 확인 방법: {t['verify_with']}"]
    md += ["", "## 11. 브랜드 표기", ""]
    for k, e in facts["brand"].items():
        md.append(f"- {k}: `{e['value']}` ({mark[e['status']]}) — 원본에서 본 표기: {', '.join(e['variants_seen'])}" + (f". {e['note']}" if e.get("note") else ""))
    return "\n".join(md) + "\n"


def main(argv=None) -> int:
    try:
        from console_encoding import configure_utf8_stdio
        configure_utf8_stdio()
    except ImportError:
        pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--pptx", type=Path)
    g.add_argument("--svg", type=Path, help="skill project folder (svg_output/) or a folder of SVGs")
    g.add_argument("--outline", type=Path)
    g.add_argument("--print-md", action="store_true")
    ap.add_argument("--facts", type=Path, default=None)
    ap.add_argument("-o", "--out", type=Path, default=None, help="with --print-md: write UTF-8 file instead of stdout")
    ap.add_argument("--today", default=None, help="YYYY-MM-DD (for review_by checks)")
    args = ap.parse_args(argv)
    facts = load_facts(args.facts)
    if facts is None:
        print(f"[facts_check] ledger not found: {args.facts or FACTS_PATH}")
        return 2
    if args.print_md:
        if args.out:
            args.out.write_text(to_markdown(facts), encoding="utf-8")
            print(f"[facts_check] wrote {args.out}")
        else:
            sys.stdout.write(to_markdown(facts))
        return 0
    pages = (pages_from_pptx(args.pptx) if args.pptx else
             pages_from_svg(args.svg) if args.svg else pages_from_outline(args.outline))
    found: list[tuple] = []
    today = _dt.date.fromisoformat(args.today) if args.today else None
    check(pages, facts, lambda *a: found.append(a), today)
    for lvl, code, page, msg in sorted(found, key=lambda x: ({"ERROR": 0, "WARN": 1, "INFO": 2}[x[0]], x[2])):
        print(f"[{lvl}] {code} {'P%02d' % page if page else '덱'} {msg}")
    warns = sum(1 for f in found if f[0] == "WARN")
    print(f"[facts_check] {len(pages)} pages · WARN {warns} · INFO {len(found) - warns}")
    return 1 if warns else 0


if __name__ == "__main__":
    raise SystemExit(main())
