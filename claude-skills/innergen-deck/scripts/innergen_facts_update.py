#!/usr/bin/env python3
"""
INNerGEN company-facts updater - writes a copy of an existing PowerPoint in
which stale company values (the F01/F02 findings of innergen_facts_check.py)
are replaced with the ledger values, keeping every run's font, size and color.

Only values whose ledger entry is `confirmed` are changed by default; values
marked `needs_confirm` (for example 감축실적 250만 vs 300만 톤, the INNerCarry
spelling) are listed but left alone unless --include-unconfirmed is given.
The input file is never overwritten.

Usage:
    python3 scripts/innergen_facts_update.py deck.pptx -o deck_updated.pptx
    python3 scripts/innergen_facts_update.py deck.pptx -o out.pptx --include-unconfirmed
    python3 scripts/innergen_facts_update.py deck.pptx --dry-run      # list changes only

Dependencies:
    python-pptx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

import innergen_facts_check as facts_check  # noqa: E402


def _entry(facts: dict, dotted: str) -> dict:
    """`metrics.credits_recognized` -> the ledger entry (lists are matched by id)."""
    node = facts
    for part in dotted.split("."):
        if isinstance(node, list):
            node = next((x for x in node if isinstance(x, dict) and x.get("id") == part), {})
        elif isinstance(node, dict):
            node = node.get(part, {})
    return node if isinstance(node, dict) else {}


def _rules(facts: dict, include_unconfirmed: bool) -> tuple[list[dict], list[str]]:
    """Replacement rules from the ledger `checks` and brand spellings."""
    rules, skipped = [], []
    for c in facts.get("checks", []):
        status = _entry(facts, c.get("fact", "")).get("status", "confirmed")
        rule = {"label": c["label"], "rx": re.compile(c["rx"]), "status": status,
                "page_rx": re.compile(c["page_rx"]) if c.get("page_rx") else None,
                "expect": c["expect"], "digits": c.get("normalize") == "digits", "kind": "value"}
        if status == "confirmed" or include_unconfirmed:
            rules.append(rule)
        else:
            skipped.append(c["label"])
    for key, pattern in facts.get("brand_rx", {}).items():
        e = facts.get("brand", {}).get(f"{key}_spelling", {})
        if not e.get("value"):
            continue
        rule = {"label": f"브랜드 표기({e['value']})", "rx": re.compile(pattern, re.IGNORECASE), "status": e.get("status"),
                "page_rx": None, "expect": [e["value"]], "digits": False, "kind": "brand"}
        if e.get("status") == "confirmed" or include_unconfirmed:
            rules.append(rule)
        else:
            skipped.append(rule["label"])
    return rules, skipped


def _replacement(rule: dict, m: re.Match) -> tuple[int, int, str] | None:
    """(start, end, new text) for the part of the match to change, or None when
    the value is already current. Only the value group is replaced, so units and
    labels keep their own runs (and font sizes)."""
    if rule["kind"] == "brand":
        got = m.group(0)
        if got.islower() or re.sub(r"\s", "", got) == rule["expect"][0]:
            return None
        return m.start(), m.end(), rule["expect"][0]
    gi = next((i for i, g in enumerate(m.groups(), 1) if g), 0)
    got = m.group(gi)
    norm = facts_check._digits if rule["digits"] else (lambda s: s.strip())
    if norm(got) in {norm(e) for e in rule["expect"]}:
        return None
    return m.start(gi), m.end(gi), rule["expect"][0]


def _apply_paragraph(par, rule: dict, log: list, slide_no: int) -> None:
    runs = list(par.runs)
    if not runs:
        return
    text = "".join(r.text for r in runs)
    edits = []
    for m in rule["rx"].finditer(text):
        rep = _replacement(rule, m)
        if rep is not None:
            start, end, new = rep
            shown = m.group(0)[: start - m.start()] + new + m.group(0)[end - m.start():]
            edits.append((start, end, new, m.group(0), shown))
    for start, end, new, old, shown in reversed(edits):  # right to left keeps offsets valid
        pos, first = 0, True
        for r in runs:
            r_start, r_end = pos, pos + len(r.text)
            pos = r_end
            if r_end <= start or r_start >= end:
                continue
            a, b = max(start, r_start) - r_start, min(end, r_end) - r_start
            r.text = r.text[:a] + (new if first else "") + r.text[b:]
            first = False
        log.append((slide_no, rule["label"], old, shown, rule["status"]))


def _text_frames(shapes):
    for s in shapes:
        if s.shape_type == 6:  # group
            yield from _text_frames(s.shapes)
            continue
        if s.has_text_frame:
            yield s.text_frame
        if getattr(s, "has_table", False) and s.has_table:
            for row in s.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def update(src: Path, out: Path | None, facts: dict, include_unconfirmed: bool) -> tuple[list, list]:
    from pptx import Presentation

    prs = Presentation(str(src))
    rules, skipped = _rules(facts, include_unconfirmed)
    log: list = []
    for no, slide in enumerate(prs.slides, 1):
        frames = list(_text_frames(slide.shapes))
        blob = "\n".join(tf.text for tf in frames)
        for rule in rules:
            if rule["page_rx"] and not rule["page_rx"].search(blob):
                continue
            for tf in frames:
                for par in tf.paragraphs:
                    _apply_paragraph(par, rule, log, no)
    if out is not None:
        prs.save(str(out))
    return log, skipped


def main(argv=None) -> int:
    try:
        from console_encoding import configure_utf8_stdio
        configure_utf8_stdio()
    except ImportError:
        pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=None, help="updated copy (required unless --dry-run)")
    ap.add_argument("--include-unconfirmed", action="store_true", help="also apply ledger values marked needs_confirm")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--facts", type=Path, default=None)
    args = ap.parse_args(argv)
    if not args.dry_run and args.out is None:
        ap.error("-o/--out is required (the input file is never overwritten)")
    if args.out is not None and args.out.resolve() == args.pptx.resolve():
        ap.error("--out must differ from the input file")
    facts = facts_check.load_facts(args.facts)
    if facts is None:
        print(f"[facts_update] ledger not found: {args.facts or facts_check.FACTS_PATH}")
        return 2
    log, skipped = update(args.pptx, None if args.dry_run else args.out, facts, args.include_unconfirmed)
    for no, label, old, new, status in log:
        mark = "" if status == "confirmed" else f" ({status})"
        print(f"[CHANGE] P{no:02d} {label}: '{old}' → '{new}'{mark}")
    if skipped:
        print(f"[facts_update] 확인 필요라서 건너뛴 항목: {', '.join(skipped)} (--include-unconfirmed로 적용)")
    target = "변경 없음(dry-run)" if args.dry_run else str(args.out)
    print(f"[facts_update] {len(log)} change(s) · {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
