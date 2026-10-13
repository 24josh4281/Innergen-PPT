#!/usr/bin/env python3
"""
PPT Master - Source Design Evidence

Read-only OOXML design observations with declared style inheritance. Document
text and instructions are not executed or included in the evidence payload.

Usage:
    Imported by audit_design_sources.py; no standalone file writes.

Dependencies:
    python-pptx, lxml
"""
from __future__ import annotations
import argparse
from collections import Counter, defaultdict
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from lxml import etree
from pptx import Presentation

NS = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
EMU=914400
_xpath_cache={}

def xp(el, path):
    if el is None:return []
    if path not in _xpath_cache:_xpath_cache[path]=etree.XPath(path,namespaces=NS)
    return _xpath_cache[path](el)

def first(el,path):
    found=xp(el,path)
    return found[0] if found else None

def digest(path):
    h=hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1048576),b''): h.update(chunk)
    return h.hexdigest()

def rgb_transform(rgb, node):
    if not rgb or node is None:return rgb
    try: channels=[int(rgb[i:i+2],16)/255 for i in (0,2,4)]
    except ValueError:return rgb
    for child in node:
        action=etree.QName(child).localname
        val=float(child.get('val','0'))/100000
        if action=='tint':channels=[c+(1-c)*val for c in channels]
        elif action=='shade':channels=[c*val for c in channels]
        elif action=='lumMod':channels=[c*val for c in channels]
        elif action=='lumOff':channels=[c+val for c in channels]
    return ''.join(f'{max(0,min(255,round(c*255))):02X}' for c in channels)

def get_theme(slide):
    master=slide.slide_layout.slide_master
    rels=[r for r in master.part.rels.values() if r.reltype.endswith('/theme')]
    theme=etree.fromstring(rels[0].target_part.blob) if rels else None
    cmap=first(master._element,'./p:clrMap')
    override=first(slide._element,'./p:clrMapOvr/a:overrideClrMapping')
    mapping=dict((override if override is not None else cmap).attrib) if (override is not None or cmap is not None) else {}
    colors={}
    scheme=first(theme,'./a:themeElements/a:clrScheme')
    if scheme is not None:
        for item in scheme:
            c=item[0] if len(item) else None
            if c is not None:colors[etree.QName(item).localname]=c.get('val') if etree.QName(c).localname=='srgbClr' else c.get('lastClr')
    fonts={}
    for category in ('major','minor'):
        el=first(theme,f'./a:themeElements/a:fontScheme/a:{category}Font')
        fonts[category]={}
        if el is not None:
            for child in el:
                kind=etree.QName(child).localname
                fonts[category][child.get('script',kind)]=child.get('typeface','')
    return {'theme':theme,'colors':colors,'map':mapping,'fonts':fonts,'master':master}

def resolve_color(node,theme):
    if node is None:return None
    rgb=first(node,'./a:srgbClr')
    if rgb is not None:return {'rgb':rgb_transform(rgb.get('val'),rgb),'type':'rgb','base':rgb.get('val')}
    scheme=first(node,'./a:schemeClr')
    if scheme is not None:
        key=scheme.get('val'); mapped=theme['map'].get(key,key)
        return {'rgb':rgb_transform(theme['colors'].get(mapped),scheme),'type':'theme','scheme':key,'mapped_scheme':mapped}
    system=first(node,'./a:sysClr')
    if system is not None:return {'rgb':system.get('lastClr'),'type':'system'}
    return None

def theme_font(raw,theme,kind):
    if not raw:return None
    if not raw.startswith('+'):return raw
    family='major' if raw.startswith('+mj') else 'minor'
    category={'lt':'latin','ea':'ea','cs':'cs'}.get(raw[-2:],kind)
    fonts=theme['fonts'].get(family,{})
    if category=='ea':return fonts.get('ea') or fonts.get('Hang') or None
    return fonts.get(category) or None

def ph_info(el):
    node=first(el,'./p:nvSpPr/p:nvPr/p:ph')
    return {'idx':node.get('idx','0'),'type':node.get('type','obj')} if node is not None else None

def matching_placeholder(container,ph):
    if ph is None:return None
    candidates=xp(container,'./p:cSld/p:spTree/p:sp')
    for candidate in candidates:
        other=ph_info(candidate)
        if other and other['idx']==ph['idx']:return candidate
    for candidate in candidates:
        other=ph_info(candidate)
        if other and other['type']==ph['type']:return candidate
    return None

def text_chain(shape_el,p,r,slide,level):
    result=[]
    for name,node in [('run',first(r,'./a:rPr')),('paragraph',first(p,'./a:pPr/a:defRPr')),
                      ('shape',first(shape_el,f'./p:txBody/a:lstStyle/a:lvl{level+1}pPr/a:defRPr'))]:
        if node is not None:result.append((name,node))
    ph=ph_info(shape_el)
    layout_shape=matching_placeholder(slide.slide_layout._element,ph)
    master_shape=matching_placeholder(slide.slide_layout.slide_master._element,ph)
    for prefix,source in [('layout',layout_shape),('master-placeholder',master_shape)]:
        if source is None:continue
        for path in [f'./p:txBody/a:lstStyle/a:lvl{level+1}pPr/a:defRPr','./p:txBody/a:p/a:pPr/a:defRPr','./p:txBody/a:p/a:r/a:rPr']:
            node=first(source,path)
            if node is not None:result.append((prefix,node))
    title=ph and ph['type'] in ('title','ctrTitle')
    body=ph and ph['type'] in ('body','obj','subTitle')
    style='titleStyle' if title else 'bodyStyle' if body else 'otherStyle'
    node=first(slide.slide_layout.slide_master._element,f'./p:txStyles/p:{style}/a:lvl{level+1}pPr/a:defRPr')
    if node is not None:result.append(('master-text-style',node))
    return result

def text_property(chain,prop,theme,kind=None):
    for source,node in chain:
        if prop in ('size','bold','italic'):
            raw=node.get({'size':'sz','bold':'b','italic':'i'}[prop])
            if raw is not None:return {'value':float(raw)/100 if prop=='size' else raw=='1','source':source,'explicit':source=='run'}
        elif prop=='font':
            child=first(node,f'./a:{kind}')
            if child is not None:
                raw=child.get('typeface'); resolved=theme_font(raw,theme,kind)
                if resolved:return {'value':resolved,'raw':raw,'source':source,'explicit':source=='run','theme_reference':bool(raw and raw.startswith('+'))}
        elif prop=='color':
            fill=first(node,'./a:solidFill')
            if fill is not None:
                found=resolve_color(fill,theme)
                if found:return {**found,'source':source,'explicit':source=='run'}
    return {'value':None,'source':'unresolved','explicit':False} if prop!='color' else {'rgb':None,'source':'unresolved','explicit':False}

def shape_paint(shape_el,slide,theme,stroke=False):
    tail='/a:ln' if stroke else ''
    prop=first(shape_el,'./p:spPr'+tail)
    if prop is not None:
        if first(prop,'./a:noFill') is not None:return {'type':'none','source':'shape','explicit':True}
        solid=first(prop,'./a:solidFill')
        if solid is not None:return {**(resolve_color(solid,theme) or {'type':'unresolved'}),'source':'shape','explicit':True}
        if first(prop,'./a:gradFill') is not None:return {'type':'gradient','source':'shape','explicit':True}
    ph=ph_info(shape_el)
    for label,container in [('layout',slide.slide_layout._element),('master',slide.slide_layout.slide_master._element)]:
        inherited=matching_placeholder(container,ph)
        inherited_prop=first(inherited,'./p:spPr'+tail)
        if inherited_prop is not None:
            if first(inherited_prop,'./a:noFill') is not None:return {'type':'none','source':label,'explicit':False}
            solid=first(inherited_prop,'./a:solidFill')
            if solid is not None:return {**(resolve_color(solid,theme) or {'type':'unresolved'}),'source':label,'explicit':False}
    style=first(shape_el,'./p:style/a:lnRef' if stroke else './p:style/a:fillRef')
    if style is not None:
        if style.get('idx')=='0':return {'type':'none','source':'style-ref','explicit':False,'style_index':'0'}
        found=resolve_color(style,theme)
        if found:return {**found,'source':'style-ref','explicit':False,'style_index':style.get('idx')}
    return {'type':'unresolved','source':'unresolved','explicit':False}

def rect(shape,w,h):
    return {'x_in':round(shape.left/EMU,3),'y_in':round(shape.top/EMU,3),'w_in':round(shape.width/EMU,3),'h_in':round(shape.height/EMU,3),
            'x_ratio':round(shape.left/w,4),'y_ratio':round(shape.top/h,4),'w_ratio':round(shape.width/w,4),'h_ratio':round(shape.height/h,4)}

def walk_shapes(shapes,group=False):
    for shape in shapes:
        yield shape,group
        if shape.shape_type==6:
            yield from walk_shapes(shape.shapes,True)

def role_for(shape,runs,geometry,w,h):
    ph=ph_info(shape._element)
    if ph and ph['type'] in ('title','ctrTitle'):return 'title'
    if getattr(shape,'has_table',False):return 'table'
    if getattr(shape,'has_chart',False):return 'chart'
    if shape.shape_type==13:return 'image'
    if shape.shape_type==6:return 'group'
    if ph and ph['type'] in ('ftr','dt','sldNum'):return 'footer'
    if geometry['y_ratio']>=.9:return 'footer'
    if geometry['y_ratio']<=.15 and runs:return 'title-region'
    sizes=[r['size']['value'] for r in runs if r['size']['value'] is not None]
    if runs:
        if sizes and max(sizes)>=20 and geometry['h_ratio']<=.14:return 'section-heading'
        return 'body-text'
    if geometry['h_ratio']<=.035 or geometry['w_ratio']<=.01:return 'rule-or-line'
    return 'panel-or-diagram'

def analyze(path,root):
    prs=Presentation(path)
    themes={}
    templates={}
    deck={'source':str(path.relative_to(root)).replace('\\','/'),'sha256':digest(path),'file_bytes':path.stat().st_size,
          'slides':len(prs.slides),'size_in':[round(prs.slide_width/EMU,4),round(prs.slide_height/EMU,4)],'slide_evidence':[]}
    fonts=Counter(); sizes=Counter(); textcolors=Counter(); fills=Counter(); lines=Counter(); role_positions=defaultdict(Counter)
    per_role_fonts=defaultdict(Counter); placeholder_count=Counter(); unresolved=Counter(); source_counts=Counter()
    for slide_index,slide in enumerate(prs.slides,1):
        theme=get_theme(slide); theme_key=slide.slide_layout.slide_master.part.partname
        themes[str(theme_key)]={'name':theme['theme'].get('name') if theme['theme'] is not None else None,'colors':theme['colors'],'fonts':theme['fonts']}
        evidence={'slide':slide_index,'layout_name':slide.slide_layout.name,'master_name':slide.slide_layout.slide_master.name,
                  'shape_count':0,'text_character_count':0,'table_count':0,'chart_count':0,'picture_count':0,'shapes':[]}
        all_sizes=[]; roles=Counter()
        for shape,grouped in walk_shapes(slide.shapes):
            geometry=rect(shape,prs.slide_width,prs.slide_height)
            rawruns=[]
            for p in (xp(shape._element,'.//a:p') if shape.shape_type!=6 else []):
                level=int(first(p,'./a:pPr').get('lvl','0')) if first(p,'./a:pPr') is not None else 0
                for r in xp(p,'./a:r | ./a:fld'):
                    t=first(r,'./a:t'); text=t.text or '' if t is not None else ''
                    if not text.strip():continue
                    chain=text_chain(shape._element,p,r,slide,level)
                    latin=text_property(chain,'font',theme,'latin'); ea=text_property(chain,'font',theme,'ea')
                    hangul=bool(re.search('[가-힣]',text))
                    effective=ea if hangul and ea['value'] else latin
                    effective=dict(effective); effective['choice']='east-asian' if hangul and ea['value'] else 'latin'
                    run={'characters':len(text),'contains_hangul':hangul,'size':text_property(chain,'size',theme),
                         'font_latin':latin,'font_east_asian':ea,'font_effective':effective,'bold':text_property(chain,'bold',theme),
                         'color':text_property(chain,'color',theme),'paragraph_level':level}
                    rawruns.append(run)
            role=role_for(shape,rawruns,geometry,prs.slide_width,prs.slide_height)
            roles[role]+=1
            fill=shape_paint(shape._element,slide,theme)
            line=shape_paint(shape._element,slide,theme,True)
            shape_evidence={'id':shape.shape_id,'kind':str(shape.shape_type),'role':role,'geometry':geometry,'group_child':grouped,
                            'placeholder':ph_info(shape._element),'fill':fill,'line':line,'text_runs':rawruns}
            if shape_evidence['placeholder']:placeholder_count[shape_evidence['placeholder']['type']]+=1
            if getattr(shape,'has_table',False):
                shape_evidence['table_dimensions']=[len(shape.table.rows),len(shape.table.columns)]
                evidence['table_count']+=1
            if getattr(shape,'has_chart',False):evidence['chart_count']+=1
            if shape.shape_type==13:evidence['picture_count']+=1
            evidence['shape_count']+=1
            evidence['text_character_count']+=sum(r['characters'] for r in rawruns)
            evidence['shapes'].append(shape_evidence)
            if not grouped and role in ('title','title-region','section-heading','body-text','footer'):
                role_positions[role][(round(geometry['x_in'],1),round(geometry['y_in'],1),round(geometry['w_in'],1),round(geometry['h_in'],1))]+=1
            for paint,target in [(fill,fills),(line,lines)]:
                if paint.get('rgb'):target[(role,paint['rgb'],paint['source'],str(paint['explicit']))]+=1
            for run in rawruns:
                eff=run['font_effective']; size=run['size']; color=run['color']
                fonts[(eff['value'],eff['source'],str(eff['explicit']),eff.get('raw'),eff.get('choice'))]+=1
                per_role_fonts[role][(eff['value'],size['value'])]+=1
                sizes[(role,size['value'],size['source'],str(size['explicit']))]+=1
                if size['value'] is not None:all_sizes.append(size['value'])
                if color.get('rgb'):textcolors[(role,color['rgb'],color['source'],str(color['explicit']))]+=1
                for label,prop in [('font',eff),('size',size),('text-color',color)]:
                    source_counts[(label,prop['source'])]+=1
                    if prop['source']=='unresolved':unresolved[label]+=1
        evidence['role_counts']=dict(roles)
        evidence['largest_font_pt']=max(all_sizes) if all_sizes else None
        deck['slide_evidence'].append(evidence)
        for label,container in [('master',slide.slide_layout.slide_master),('layout',slide.slide_layout)]:
            part=str(container.part.partname)
            if part in templates:
                templates[part]['referenced_slide_count']+=1
                continue
            entry={'kind':label,'part':part,'name':container.name,'referenced_slide_count':1,'shapes':[]}
            bg=first(container._element,'./p:cSld/p:bg/p:bgPr/a:solidFill')
            entry['background']=resolve_color(bg,theme)
            for ts,tgroup in walk_shapes(container.shapes):
                tg=rect(ts,prs.slide_width,prs.slide_height)
                tfill=shape_paint(ts._element,slide,theme)
                tline=shape_paint(ts._element,slide,theme,True)
                tchar=sum(len(t.text or '') for t in xp(ts._element,'.//a:t')) if ts.shape_type!=6 else 0
                trole='template-picture' if ts.shape_type==13 else 'template-text' if tchar else 'template-decoration'
                if tg['w_ratio']>.65 and tg['h_ratio']<.14 and (tg['y_ratio']<.15 or tg['y_ratio']>.85):trole='brand-band-candidate'
                tentry={'shape_id':ts.shape_id,'kind':str(ts.shape_type),'role':trole,'geometry':tg,'group_child':tgroup,
                        'fill':tfill,'line':tline,'placeholder':ph_info(ts._element),'text_character_count':tchar}
                if ts.shape_type==13:
                    try:
                        tentry['image_sha256']=hashlib.sha256(ts.image.blob).hexdigest()
                        tentry['image_size']=ts.image.size
                    except ValueError:
                        tentry['image_status']='No embedded image; external linked images were not fetched.'
                entry['shapes'].append(tentry)
            templates[part]=entry
    deck['themes']=themes
    deck['design_templates']=list(templates.values())
    deck['stats']={'font_runs':counter_rows(fonts,['font','source','explicit','raw','choice']),
                   'size_runs':counter_rows(sizes,['role','size_pt','source','explicit']),
                   'text_colors':counter_rows(textcolors,['role','rgb','source','explicit']),
                   'shape_fills':counter_rows(fills,['role','rgb','source','explicit']),
                   'shape_lines':counter_rows(lines,['role','rgb','source','explicit']),
                   'role_font_sizes':{role:counter_rows(values,['font','size_pt']) for role,values in per_role_fonts.items()},
                   'role_positions':{role:counter_rows(values,['x_in','y_in','w_in','h_in']) for role,values in role_positions.items()},
                   'placeholder_types':dict(placeholder_count),'unresolved_run_properties':dict(unresolved),
                   'property_provenance':counter_rows(source_counts,['property','source'])}
    return deck

def counter_rows(counter,fields):
    return [dict(zip(fields,key if isinstance(key,tuple) else [key]),count=n) for key,n in counter.most_common()]

def aggregate(decks):
    output={}
    for field in ['font_runs','size_runs','text_colors','shape_fills','shape_lines','property_provenance']:
        counts=Counter()
        for deck in decks:
            for row in deck['stats'][field]:counts[tuple((k,v) for k,v in row.items() if k!='count')]+=row['count']
        output[field]=[dict(key,count=count) for key,count in counts.most_common()]
    for field in ['role_positions','role_font_sizes']:
        groups=defaultdict(Counter)
        for deck in decks:
            for role,rows in deck['stats'][field].items():
                for row in rows:groups[role][tuple((k,v) for k,v in row.items() if k!='count')]+=row['count']
        output[field]={role:[dict(key,count=count) for key,count in counts.most_common()] for role,counts in groups.items()}
    return output

def choose_representatives(deck):
    slides=deck['slide_evidence']
    reps=[{'role':'cover','slide':1}]
    bodies=[s for s in slides[1:] if s['text_character_count']>100 and s['table_count']==0]
    if bodies:
        target=sorted(bodies,key=lambda s:s['text_character_count'])[len(bodies)//2]
        reps.append({'role':'typical-body','slide':target['slide']})
    tables=[s for s in slides if s['table_count']]
    if tables:reps.append({'role':'table','slide':max(tables,key=lambda s:s['table_count'])['slide']})
    diagrams=[s for s in slides[1:] if s['role_counts'].get('panel-or-diagram',0)>=6 and s['text_character_count']>50]
    if diagrams:reps.append({'role':'process-or-diagram-candidate','slide':max(diagrams,key=lambda s:s['role_counts'].get('panel-or-diagram',0))['slide']})
    reps.append({'role':'closing-candidate','slide':len(slides)})
    return reps
