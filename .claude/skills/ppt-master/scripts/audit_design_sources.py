#!/usr/bin/env python3
"""
PPT Master - Input Design Audit

Audit every input PPTX and additional named sources without modifying originals.
Write traceable per-file evidence, a compact manifest and a Korean review report.
Geometry is observed, not a claim that PowerPoint rendered without clipping.

Usage:
    python scripts/audit_design_sources.py --source reference.pptx

Dependencies:
    python-pptx, lxml
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation

from pptx_design_evidence import analyze, digest, walk_shapes


def aggregate(rows, keys, limit=12):
    counts = Counter()
    for row in rows:
        counts[tuple(row.get(k) for k in keys)] += row['count']
    return [dict(zip(keys, key), count=n) for key, n in counts.most_common(limit)]


def native_settings(path):
    prs = Presentation(path)
    counts, padding, lines, align, spacing = (Counter() for _ in range(5))
    outside = []
    for index, slide in enumerate(prs.slides, 1):
        if slide._element.get('show') == '0':
            counts['hidden_slides'] += 1
        for shape, grouped in walk_shapes(slide.shapes):
            el = shape._element
            for ln in el.xpath('./p:spPr/a:ln[@w]'):
                lines[round(int(ln.get('w')) / 12700, 3)] += 1
            if not getattr(shape, 'has_text_frame', False) or not shape.text.strip():
                continue
            counts['text_frames'] += 1
            body = shape.text_frame._txBody.bodyPr
            padding[tuple(round(v / 12700, 3) for v in (
                shape.text_frame.margin_left, shape.text_frame.margin_top,
                shape.text_frame.margin_right, shape.text_frame.margin_bottom))] += 1
            for child in body:
                tag = etree.QName(child).localname
                if tag in {'spAutoFit', 'normAutofit', 'noAutofit'}:
                    counts[tag] += 1
                if tag == 'normAutofit' and int(child.get('fontScale', '100000')) < 100000:
                    counts['stored_font_shrink'] += 1
            for paragraph in el.xpath('./p:txBody/a:p/a:pPr'):
                align[paragraph.get('algn', 'unspecified')] += 1
                for line in paragraph.xpath('./a:lnSpc/*'):
                    spacing[(etree.QName(line).localname, line.get('val'))] += 1
            if grouped or shape.rotation:
                continue
            x, y, w, h = shape.left, shape.top, shape.width, shape.height
            spill = max(-x, -y, x + w - prs.slide_width, y + h - prs.slide_height) / 12700
            if spill > 1:
                outside.append({'slide': index, 'shape_id': shape.shape_id,
                                'spill_pt': round(spill, 3),
                                'box_pt': [round(v / 12700, 3) for v in (x, y, w, h)]})
    return {'canvas_pt': [prs.slide_width / 12700, prs.slide_height / 12700],
            'settings': dict(counts), 'insets_ltrb_pt': padding.most_common(),
            'line_widths_pt': lines.most_common(), 'paragraph_alignment': align.most_common(),
            'line_spacing_ooxml': spacing.most_common(), 'offcanvas_candidates': outside}


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def write_readable_rules(output, records):
    """Expose per-file metrics as readable tables; leave source text out."""
    import csv
    rule_dir = output / 'rules'
    rule_dir.mkdir(exist_ok=True)
    with (output/'shape-geometry-pt.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(['file_id','slide','shape_id','role','coordinate_scope','x_pt','y_pt','width_pt','height_pt'])
        for rec in records:
            evidence = json.loads((output/rec['evidence']).read_text(encoding='utf-8'))
            stats = evidence['stats']
            md = [f"# {rec['id']} · {Path(rec['source']).name}", '',
                  '원본 관측 수치다. 새 제작값은 통합 v17 기준을 따른다. 원본은 수정하지 않았다.', '',
                  f"{rec['slides']}장 · 화면 {rec['canvas_pt'][0]:.3f} × {rec['canvas_pt'][1]:.3f}pt", '',
                  '## 글꼴', '', '| 글꼴 | 글자 구간 수 |','|---|---:|']
            md += [f"| {r['font']} | {r['count']} |" for r in aggregate(stats['font_runs'],['font'],1000)]
            md += ['', '## 역할별 글자 크기', '', '역할은 위치 기반 추정이다. 구간 수는 글자 수가 아니다.', '',
                   '| 추정 역할 | 크기 pt | 구간 수 |','|---|---:|---:|']
            md += [f"| {r['role']} | {r['size_pt']} | {r['count']} |" for r in aggregate(stats['size_runs'],['role','size_pt'],10000)]
            for heading,key in [('글자색','text_colors'),('도형 채움색','shape_fills')]:
                md += ['',f'## {heading}', '', '| HEX | 구간/도형 수 |','|---|---:|']
                md += [f"| #{r['rgb']} | {r['count']} |" for r in aggregate(stats[key],['rgb'],10000)]
            md += ['', '## 상자 안쪽 여백', '', '| 왼쪽 pt | 위 pt | 오른쪽 pt | 아래 pt | 수 |','|---:|---:|---:|---:|---:|']
            md += ['| '+' | '.join(str(v) for v in list(padding)+[count])+' |' for padding,count in rec['native']['insets_ltrb_pt']]
            md += ['', '## 선·문단 설정', '', '| 선 두께 pt | 수 |','|---:|---:|']
            md += [f'| {width} | {count} |' for width,count in rec['native']['line_widths_pt']]
            md += ['', '문단 정렬(원본 OOXML): '+str(rec['native']['paragraph_alignment']), '',
                   '줄간격(원본 OOXML): '+str(rec['native']['line_spacing_ooxml']), '',
                   '`spcPct`: 100000=100%. `spcPts`: 100=1pt. unspecified는 직접 지정값 없음이며 상속 확인이 필요하다.', '',
                   '자동 맞춤 등: '+str(rec['native']['settings']), '',
                   '## 대표 배치 좌표', '', '좌표는 pt. 같은 역할의 자주 쓰인 상자 최대 8개. 그룹 로컬 좌표는 아래 대표값에서 제외했다. 전체 좌표는 CSV와 원본 증거에 있다.', '',
                   '| 추정 역할 | x | y | 너비 | 높이 | 수 |','|---|---:|---:|---:|---:|---:|']
            positions = Counter()
            for page in evidence['slide_evidence']:
                for shape in page['shapes']:
                    box=shape.get('geometry') or {}
                    if not all(k in box for k in ('x_in','y_in','w_in','h_in')): continue
                    rect=tuple(round(box[k]*72,3) for k in ('x_in','y_in','w_in','h_in'))
                    grouped=shape.get('group_child',False)
                    writer.writerow([rec['id'],page['slide'],shape['id'],shape['role'],'group_local' if grouped else 'slide',*rect])
                    if not grouped:positions[(shape['role'],*rect)]+=1
            roles=sorted({r[0] for r in positions})
            for role in roles:
                rows=[(key,count) for key,count in positions.most_common() if key[0]==role][:8]
                md += ['| '+' | '.join(str(v) for v in list(key)+[count])+' |' for key,count in rows]
            md += ['', '## 화면 밖 상자 후보', '', '| 장표 | 도형 ID | 초과 pt | x/y/너비/높이 pt |','|---:|---:|---:|---|']
            md += [f"| {r['slide']} | {r['shape_id']} | {r['spill_pt']} | {r['box_pt']} |" for r in rec['native']['offcanvas_candidates']]
            md += ['', '실제 잘림 여부는 렌더링으로 확인해야 한다. 이미지 내부 글씨, 회전·그룹의 화면 경계는 이 후보 검사에 포함하지 않는다.', '',
                   f"[전체 장표·도형·글자 구간 증거](../sources/{rec['id']}.json)", '']
            (rule_dir/f"{rec['id']}.md").write_text('\n'.join(md),encoding='utf-8')


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=Path('INPUTS'))
    parser.add_argument('--source', type=Path, action='append', default=[])
    parser.add_argument('--output', type=Path, default=Path('OUTPUTS/design-audit-v17'))
    args = parser.parse_args(argv)
    paths = sorted(args.input.rglob('*.pptx')) + args.source
    paths = list(dict.fromkeys(p.resolve() for p in paths))
    if not paths or any(not p.is_file() for p in paths):
        parser.error('All requested PPTX sources must exist.')
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    records = []
    for number, path in enumerate(paths, 1):
        ident = f'S{number:02d}'
        print(f'{ident}: {path.name}', file=sys.stderr, flush=True)
        evidence = analyze(path, path.parent)
        native = native_settings(path)
        if digest(path) != evidence['sha256']:
            raise RuntimeError(f'Source changed while reading: {path}')
        evidence['absolute_source'] = str(path)
        evidence['native_settings'] = native
        evidence_file = output / 'sources' / f'{ident}.json'
        # Compact full evidence avoids hundreds of MB of repeated indentation.
        evidence_file.parent.mkdir(parents=True, exist_ok=True)
        evidence_file.write_text(json.dumps(evidence, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
        stats = evidence['stats']
        record = {'id': ident, 'source': str(path), 'sha256': evidence['sha256'],
                  'slides': evidence['slides'], 'canvas_pt': native['canvas_pt'],
                  'evidence': str(evidence_file.relative_to(output)),
                  'font_runs': aggregate(stats['font_runs'], ['font'], 30),
                  'sizes_by_role': {role: aggregate([r for r in stats['size_runs'] if r['role'] == role], ['size_pt'], 30)
                                    for role in ('title', 'title-region', 'section-heading', 'body-text', 'table', 'footer')},
                  'text_colors': aggregate(stats['text_colors'], ['rgb'], 30),
                  'shape_fills': aggregate(stats['shape_fills'], ['rgb'], 30),
                  'positions': stats['role_positions'], 'native': native,
                  'unresolved': stats['unresolved_run_properties'],
                  'tables': sum(p['table_count'] for p in evidence['slide_evidence']),
                  'charts': sum(p['chart_count'] for p in evidence['slide_evidence']),
                  'picture_instances': sum(p['picture_count'] for p in evidence['slide_evidence'])}
        records.append(record)
    manifest = {'version': '17', 'date': '2026-10-13', 'files': records,
                'totals': {'files': len(records), 'slides': sum(d['slides'] for d in records),
                           'hidden_slides': sum(d['native']['settings'].get('hidden_slides', 0) for d in records),
                           'unique_hashes': len({d['sha256'] for d in records})},
                'limitations': ['원본 설정·상속값이며 렌더링 후 실제 글자 크기나 잘림을 보증하지 않음.',
                                '그룹 내부 좌표는 그룹 로컬 좌표. group_child=true이며 전역 좌표로 혼동하지 않음.',
                                '역할 분류는 위치 기반 추정. 장표·도형 ID로 원본 대조 필요.',
                                '이미지 안의 글꼴·색·도표는 편집 가능한 텍스트/차트 통계에 포함하지 않음.',
                                '파일 내용과 과거 고객 수치는 새 장표의 사실로 재사용하지 않음.']}
    write_json(output / 'manifest.json', manifest)
    write_readable_rules(output, records)
    md = ['# 전체 입력 PPT 디자인 실측 목록', '', '기준일: 2026-10-13. 원본 파일은 수정하지 않았다.',
          '', f"총 {len(records)}개 파일, {manifest['totals']['slides']}장(숨김 {manifest['totals']['hidden_slides']}장 포함).",
          '', '각 파일의 세부 증거는 `sources/Sxx.json`에 장표·도형·글자 구간별로 저장한다. `manifest.json`은 전체 요약이다.', '',
          '## 읽는 방법', '', '원본의 자주 쓰인 수치와 새 제작 기준은 다르다. 아래 값은 원본 관측값이며 새 제작 기준은 `docs/styles/innergen/design-standard-v17.md`를 따른다.', '',
          '1pt = 1/72인치 = 0.352778mm. 기존 증거의 인치 좌표는 72를 곱하면 pt다. 그룹 내부 좌표를 화면 좌표로 사용하지 않는다.', '']
    for d in records:
        body = ', '.join(str(r['size_pt']) for r in d['sizes_by_role']['body-text'][:5])
        table = ', '.join(str(r['size_pt']) for r in d['sizes_by_role']['table'][:5]) or '편집 가능한 표 없음'
        md += [f"## {d['id']} · {Path(d['source']).name}", '', f"- 원본: `{d['source']}`", f"- 장수/화면: {d['slides']}장 / {d['canvas_pt'][0]} × {d['canvas_pt'][1]}pt",
               '- 글꼴: ' + ', '.join(f"{r['font']}({r['count']}구간)" for r in d['font_runs'][:5]),
               f'- 본문 크기 빈도 상위: {body}pt', f'- 표 크기 빈도 상위: {table}',
               '- 글자색: ' + ', '.join('#' + r['rgb'] for r in d['text_colors'][:8]),
               '- 도형 채움: ' + ', '.join('#' + r['rgb'] for r in d['shape_fills'][:8]),
               '- 선 두께(pt·개수): ' + str(d['native']['line_widths_pt'][:8]),
               '- 상자 안쪽 여백(좌·상·우·하 pt·개수): ' + str(d['native']['insets_ltrb_pt'][:5]),
               f"- 화면 밖 텍스트 상자 후보: {len(d['native']['offcanvas_candidates'])}개. 렌더 후 실제 잘림 여부 판단.",
               f"- 원본 검증값: `{d['sha256']}`", f"- [상세 수치표](rules/{d['id']}.md) · [전체 증거](sources/{d['id']}.json)", '']
    md += ['## 분석 한계', ''] + ['- ' + x for x in manifest['limitations']]
    (output / 'source-design-audit.md').write_text('\n'.join(md) + '\n', encoding='utf-8')
    print(json.dumps(manifest['totals'], ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
