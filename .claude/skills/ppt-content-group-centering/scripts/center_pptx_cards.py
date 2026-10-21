#!/usr/bin/env python3
"""Safely center explicitly mapped PowerPoint content blocks within card backgrounds.

Dry-run by default; --apply writes a separate editable PPTX.
This is a shape-geometry operation, not optical glyph-bound measurement.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    print("Missing dependency: pip install python-pptx", file=sys.stderr)
    raise SystemExit(2)

DEFAULT_PADDING = {"top": 0.12, "bottom": 0.12, "left": 0.18, "right": 0.18}
EMU_INCH = 914400


def to_in(value: int | float) -> float:
    return round(float(value) / EMU_INCH, 4)


def find_unique_shape(slide: Any, name: str) -> Any:
    matched = [shape for shape in slide.shapes if shape.name == name]
    if len(matched) != 1:
        raise ValueError(f"shape name {name!r}: expected 1, found {len(matched)}")
    return matched[0]


def union_bounds(shapes: list[Any]) -> tuple[int, int, int, int]:
    left = min(int(s.left) for s in shapes)
    top = min(int(s.top) for s in shapes)
    right = max(int(s.left + s.width) for s in shapes)
    bottom = max(int(s.top + s.height) for s in shapes)
    return left, top, right, bottom


def process_card(pres: Any, entry: dict[str, Any], apply: bool) -> dict[str, Any]:
    slide_number = entry.get("slide_number")
    label = f"slide {slide_number}, background {entry.get('background_shape_name', '?')}"
    base = {"card": label, "status": "skipped", "warnings": [], "render_qa": "not_performed"}
    if not isinstance(slide_number, int) or not 1 <= slide_number <= len(pres.slides):
        return {**base, "reason": "invalid 1-based slide_number"}
    try:
        slide = pres.slides[slide_number - 1]
        bg = find_unique_shape(slide, entry["background_shape_name"])
        names = entry["content_shape_names"]
        if not isinstance(names, list) or not names or len(names) != len(set(names)):
            raise ValueError("content_shape_names must be nonempty and unique")
        if bg.name in names:
            raise ValueError("background cannot be part of content")
        children = [find_unique_shape(slide, name) for name in names]
        if any(c.width <= 0 or c.height <= 0 for c in children):
            raise ValueError("content has zero/negative geometry")
        raw_pad = {**DEFAULT_PADDING, **entry.get("padding_inches", {})}
        if any(not isinstance(raw_pad[k], (float, int)) or raw_pad[k] < 0 for k in DEFAULT_PADDING):
            raise ValueError("padding must be nonnegative inch numbers")
        pad = {k: Inches(float(v)) for k, v in raw_pad.items()}
        left, top, right, bottom = union_bounds(children)
        b_left, b_top, b_right, b_bottom = int(bg.left), int(bg.top), int(bg.left + bg.width), int(bg.top + bg.height)
        b_h = b_bottom - b_top
        group_h = bottom - top
        if group_h > b_h - pad["top"] - pad["bottom"]:
            raise ValueError(f"content too tall for vertical safe area: group={to_in(group_h)}in card={to_in(b_h)}in")
        if left < b_left + pad["left"] or right > b_right - pad["right"]:
            raise ValueError("content crosses horizontal safe padding; no automatic x correction")
        target_top = b_top + (b_h - group_h) // 2
        desired_delta = target_top - top
        new_top, new_bottom = top + desired_delta, bottom + desired_delta
        if new_top < b_top + pad["top"] or new_bottom > b_bottom - pad["bottom"]:
            raise ValueError("centering would violate vertical safe padding")
        warnings: list[str] = []
        for child in children:
            if getattr(child, "has_text_frame", False):
                frame = child.text_frame
                if frame.margin_top > Inches(0.08) or frame.margin_bottom > Inches(0.08):
                    warnings.append(f"{child.name}: text frame internal margins may affect optical centering")
        if apply:
            for child in children:
                child.top = int(child.top) + desired_delta
        return {
            **base,
            "status": "applied" if apply else "ready_dry_run",
            "delta_y_inches": to_in(desired_delta),
            "padding_before_inches": {"top": to_in(top - b_top), "bottom": to_in(b_bottom - bottom)},
            "padding_after_inches": {"top": to_in(new_top - b_top), "bottom": to_in(b_bottom - new_bottom)},
            "safe_padding_inches": raw_pad,
            "child_names": names,
            "warnings": warnings,
        }
    except (KeyError, TypeError, ValueError) as exc:
        return {**base, "reason": str(exc)}


def run(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Existing PPTX")
    parser.add_argument("--config", required=True, type=Path, help="Cards JSON mapping")
    parser.add_argument("--apply", action="store_true", help="Move text shapes and write new PPTX (default: dry run)")
    parser.add_argument("--output", type=Path, help="Output PPTX; default input_stem_centered.pptx")
    parser.add_argument("--report", type=Path, help="Optional JSON QA report")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacing an existing output file")
    parser.add_argument("--strict", action="store_true", help="Exit failure if any card was skipped")
    args = parser.parse_args(argv)
    if not args.input.exists() or not args.config.exists():
        parser.error("input pptx or config json not found")
    cfg = json.loads(args.config.read_text(encoding="utf-8-sig"))
    if not isinstance(cfg.get("cards"), list) or not cfg["cards"]:
        parser.error("config must have a nonempty cards array")
    out = args.output or args.input.with_name(args.input.stem + "_centered.pptx")
    if args.apply and out.resolve() == args.input.resolve():
        parser.error("must write a separate PPTX; input overwrite is not permitted")
    if args.apply and out.exists() and not args.overwrite:
        parser.error(f"output already exists: {out}; supply --overwrite only after review")
    pres = Presentation(str(args.input))
    results = [process_card(pres, entry, args.apply) for entry in cfg["cards"]]
    skipped = [r for r in results if r["status"] == "skipped"]
    if args.apply and (not skipped or not args.strict):
        out.parent.mkdir(parents=True, exist_ok=True)
        pres.save(str(out))
    report = {
        "source": str(args.input),
        "output": str(out) if args.apply and (not skipped or not args.strict) else None,
        "mode": "apply" if args.apply else "dry_run",
        "total_cards": len(results),
        "skipped_cards": len(skipped),
        "caveat": "Uses PPTX shape bounding boxes, not rendered glyph bounds; render and visually QA the PPTX.",
        "results": results,
    }
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    print(payload)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(payload + "\n", encoding="utf-8")
    return 1 if args.strict and skipped else 0


if __name__ == "__main__":
    raise SystemExit(run())
