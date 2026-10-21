#!/usr/bin/env python3
"""List top-level PPTX shape names and geometry to build an explicit card mapping JSON."""
import argparse
from pptx import Presentation


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('pptx', help='path to PowerPoint file')
    ap.add_argument('--slide', type=int, help='1-based slide number; omit for all')
    args = ap.parse_args()
    pres = Presentation(args.pptx)
    if args.slide and not 1 <= args.slide <= len(pres.slides):
        ap.error('invalid slide number')
    indices = [args.slide - 1] if args.slide else range(len(pres.slides))
    for i in indices:
        print(f'--- Slide {i+1} ---')
        for sh in pres.slides[i].shapes:
            text = sh.text[:45].replace('\n', ' / ') if getattr(sh, 'has_text_frame', False) else ''
            print(f'id={sh.shape_id:>4}  name={sh.name!r:35} x={sh.left/914400:5.2f}in y={sh.top/914400:5.2f}in w={sh.width/914400:5.2f}in h={sh.height/914400:5.2f}in {text}')
        print('Note: shapes nested in a PowerPoint Group may require manual inspection.')

if __name__ == '__main__':
    main()
